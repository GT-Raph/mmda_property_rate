from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from admins.models import Assembly, UserProfile
from admins.views import PortalLoginView, _find_citizen_user_by_phone
from billing.models import Bill, CommunityArea
from payments.models import Payment
from properties.models import Owner, Property

User = get_user_model()


@override_settings(
    SECURE_SSL_REDIRECT=False,
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
)
class AdminPasswordVisibilityTests(TestCase):
    def test_admin_login_loads_password_visibility_control(self):
        response = self.client.get(reverse("admin:login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "css/password_visibility.css",
        )
        self.assertContains(
            response,
            "js/password_visibility.js",
        )


class CitizenPhoneLookupTests(TestCase):
    def setUp(self):
        self.assembly = Assembly.objects.create(
            name="Test Assembly",
            short_code="TA",
            assembly_type="municipal",
            region="Greater Accra",
            is_active=True,
        )

    def _create_citizen(self, username, phone_number=""):
        user = User.objects.create_user(
            username=username,
            password="secret123",
            is_active=True,
        )
        profile, _created = UserProfile.objects.get_or_create(
            user=user,
        )
        profile.role = "citizen"
        profile.assembly = self.assembly
        profile.phone_number = phone_number
        profile.save(
            update_fields=[
                "role",
                "assembly",
                "phone_number",
            ]
        )
        return user

    def test_find_citizen_user_by_phone_returns_single_match_once(self):
        user = self._create_citizen("citizen-one", "")
        Owner.objects.create(
            assembly=self.assembly,
            linked_user=user,
            full_name="Citizen One",
            primary_phone="233241234567",
        )

        with self.assertNumQueries(2):
            matched_user = _find_citizen_user_by_phone("0241234567")

        self.assertEqual(matched_user.pk, user.pk)


class PortalLoginViewAccessTests(TestCase):
    def test_citizen_denial_message_covers_all_access_checks(self):
        assigned_assembly = Assembly.objects.create(
            name="Assigned Assembly",
            short_code="AA",
            assembly_type="municipal",
            region="Greater Accra",
            is_active=True,
        )
        selected_assembly = Assembly.objects.create(
            name="Selected Assembly",
            short_code="SA",
            assembly_type="municipal",
            region="Greater Accra",
            is_active=True,
        )
        user = User.objects.create_user(
            username="citizen-two",
            password="secret123",
            is_active=True,
        )
        profile, _created = UserProfile.objects.get_or_create(
            user=user,
        )
        profile.role = "citizen"
        profile.assembly = assigned_assembly
        profile.phone_number = "0241234567"
        profile.save(
            update_fields=[
                "role",
                "assembly",
                "phone_number",
            ]
        )

        allowed, message = PortalLoginView().user_can_access_selected_assembly(
            user,
            selected_assembly,
        )

        self.assertFalse(allowed)
        self.assertEqual(
            str(message),
            (
                "This citizen account is not authorized to access the selected "
                "assembly because it is not linked to a property owner record, "
                "does not use the selected assembly as its default assembly, "
                "and is not included in its allowed assemblies."
            ),
        )


@override_settings(
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    }
)
class AdminAnalyticsViewTests(TestCase):
    def setUp(self):
        self.assembly = Assembly.objects.create(
            name="Analytics Assembly",
            short_code="AN",
            assembly_type="municipal",
            region="Greater Accra",
            is_active=True,
        )
        self.user = User.objects.create_superuser(
            username="analytics-admin",
            email="analytics@example.com",
            password="Admin-pass-123",
        )
        self.owner = Owner.objects.create(
            assembly=self.assembly,
            full_name="Analytics Citizen",
            primary_phone="0242222222",
        )
        self.property_record = Property.objects.create(
            assembly=self.assembly,
            owner=self.owner,
            address_line="1 Analytics Road",
            area_town="Accra",
            property_use="residential",
            building_type="apartment",
            verification_status="pending",
            is_exempt=True,
        )
        Payment.objects.create(
            assembly=self.assembly,
            property=self.property_record,
            amount=Decimal("50.00"),
            payment_method="cash",
            transaction_reference="ANALYTICS-PAYMENT-1",
            status="successful",
        )
        self.client.force_login(self.user)
        session = self.client.session
        session["active_assembly_id"] = self.assembly.pk
        session.save()

    def test_dashboard_exposes_only_four_key_metrics(self):
        response = self.client.get(
            reverse("admins:admin-dashboard"),
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        titles = [card["title"] for card in response.context["summary_cards"]]
        self.assertEqual(
            titles,
            [
                "Total Properties",
                "Pending Verification",
                "Revenue Collected",
                "Outstanding Amount",
            ],
        )
        self.assertNotContains(response, "Payment Distribution")

    def test_reports_exposes_one_filter_driven_chart_and_records(self):
        response = self.client.get(
            reverse("admins:reports-analytics"),
            {
                "report_type": "verification",
                "period": "quarterly",
                "year": timezone.localdate().year,
                "status": "awaiting_review",
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_report_type"], "verification")
        self.assertEqual(response.context["selected_period"], "quarterly")
        self.assertEqual(response.context["selected_status"], "awaiting_review")
        self.assertEqual(len(response.context["chart_config"]["labels"]), 4)
        self.assertEqual(response.context["record_title"], "Verification Records")
        self.assertEqual(response.context["record_paginator"].count, 1)
        self.assertContains(response, "Report types", html=False)
        self.assertContains(response, "Export Filtered Results (.xlsx)")

    def test_yearly_period_returns_five_year_series(self):
        response = self.client.get(
            reverse("admins:reports-analytics"),
            {"report_type": "citizens", "period": "yearly"},
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["chart_config"]["labels"]), 5)
        self.assertEqual(response.context["record_title"], "Property Owner Records")

    def test_every_report_type_renders_its_matching_records_table(self):
        expected_titles = {
            "overview": "Recent Financial Activity",
            "properties": "Property Records",
            "verification": "Verification Records",
            "billing_payments": "Billing Records",
            "citizens": "Property Owner Records",
            "disputes": "Dispute Records",
        }

        for report_type, record_title in expected_titles.items():
            with self.subTest(report_type=report_type):
                response = self.client.get(
                    reverse("admins:reports-analytics"),
                    {"report_type": report_type, "period": "monthly"},
                    secure=True,
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context["record_title"], record_title)

    def test_verification_rate_uses_the_same_filtered_period_population(self):
        verified = Property.objects.create(
            assembly=self.assembly,
            owner=self.owner,
            address_line="2 Analytics Road",
            area_town="Accra",
            property_use="residential",
            building_type="detached",
            verification_status="verified",
            is_exempt=True,
        )
        response = self.client.get(
            reverse("admins:reports-analytics"),
            {"report_type": "verification", "year": timezone.localdate().year},
            secure=True,
        )
        stats = {item["title"]: item["value"] for item in response.context["stats"]}
        self.assertEqual(stats["Properties in Scope"], "2")
        self.assertEqual(stats["Verified"], "1")
        self.assertEqual(stats["Verification Rate"], "50.0%")
        self.assertContains(response, verified.property_reference)

    def test_revenue_uses_successful_payments_only(self):
        Payment.objects.create(
            assembly=self.assembly,
            property=self.property_record,
            amount=Decimal("999.00"),
            payment_method="cash",
            transaction_reference="FAILED-PAYMENT",
            status="failed",
        )
        response = self.client.get(
            reverse("admins:reports-analytics"),
            {"report_type": "overview", "year": timezone.localdate().year},
            secure=True,
        )
        stats = {item["title"]: item["value"] for item in response.context["stats"]}
        self.assertEqual(stats["Revenue Collected"], "GH₵ 50")
        self.assertEqual(response.context["record_paginator"].count, 1)

    def test_combined_property_filters_apply_to_chart_table_and_kpis(self):
        community = CommunityArea.objects.create(
            assembly=self.assembly,
            sub_metro="AMA",
            name="Central Test Area",
        )
        self.property_record.community_area = community
        self.property_record.property_class = "Class A"
        self.property_record.save(update_fields=["community_area", "property_class", "updated_at"])
        Property.objects.create(
            assembly=self.assembly,
            owner=self.owner,
            address_line="3 Other Road",
            area_town="Accra",
            property_use="commercial",
            building_type="shop",
            verification_status="verified",
            is_exempt=True,
        )
        params = {
            "report_type": "properties",
            "year": timezone.localdate().year,
            "community_area": community.pk,
            "property_use": "residential",
            "property_class": "Class A",
            "building_type": "apartment",
            "status": "pending",
            "search": "Central Test",
        }
        response = self.client.get(reverse("admins:reports-analytics"), params, secure=True)
        self.assertEqual(response.context["record_paginator"].count, 1)
        self.assertEqual(sum(response.context["chart_config"]["datasets"][0]["data"]), 1)
        stats = {item["title"]: item["value"] for item in response.context["stats"]}
        self.assertEqual(stats["Properties Registered"], "1")
        self.assertIn("community_area=", response.context["filter_query"])

    def test_billing_outstanding_and_payment_filters_share_consistent_totals(self):
        bill = Bill.objects.get(
            property_record=self.property_record,
            billing_year=timezone.localdate().year,
        )
        Bill.objects.filter(pk=bill.pk).update(
            amount_due=Decimal("100.00"),
            amount_paid=Decimal("40.00"),
            status="part_paid",
        )
        Payment.objects.filter(transaction_reference="ANALYTICS-PAYMENT-1").update(
            bill=bill,
            amount=Decimal("40.00"),
        )
        response = self.client.get(
            reverse("admins:reports-analytics"),
            {
                "report_type": "billing_payments",
                "year": timezone.localdate().year,
                "status": "outstanding",
                "payment_activity": "recorded",
            },
            secure=True,
        )
        stats = {item["title"]: item["value"] for item in response.context["stats"]}
        self.assertEqual(response.context["record_paginator"].count, 1)
        self.assertEqual(stats["Revenue Collected"], "GH₵ 40")
        self.assertEqual(stats["Outstanding Amount"], "GH₵ 60")

    def test_filtered_export_matches_filtered_table(self):
        Property.objects.create(
            assembly=self.assembly,
            owner=self.owner,
            address_line="4 Verified Road",
            area_town="Accra",
            property_use="residential",
            building_type="detached",
            verification_status="verified",
            is_exempt=True,
        )
        params = {
            "action": "export_filtered",
            "report_type": "properties",
            "period": "monthly",
            "year": str(timezone.localdate().year),
            "status": "pending",
        }
        response = self.client.post(reverse("admins:reports-analytics"), params, secure=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        worksheet = load_workbook(BytesIO(response.content)).active
        exported = list(worksheet.iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(exported), 1)
        self.assertEqual(exported[0][1], self.property_record.property_reference)
