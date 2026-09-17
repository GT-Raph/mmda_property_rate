from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
import random

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Count, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from admins.models import Assembly, UserProfile
from billing.models import Bill, CommunityArea, RatePolicy
from billing.services import generate_bill_for_property
from field_operations.models import (
    EnumerationRecord,
    FieldAssignment,
    FieldVisit,
)
from payments.models import Payment
from properties.models import Owner, Property
from properties.utils import link_or_create_citizen_for_owner


User = get_user_model()

TARGET_ASSEMBLY_NAME = "Ablekuma South"
TARGET_ASSEMBLY_CODE = "ABS"
BILLING_YEAR = 2026
CONFIRMATION_TEXT = "REPLACE ABLEKUMA SOUTH"
PROPERTY_COUNT = 100
OWNER_COUNT = 72
VERIFIED_COUNT = 60
PENDING_COUNT = 40
FULLY_PAID_COUNT = 50
PART_PAID_COUNT = 5
UNPAID_COUNT = 5
DEMO_NOTE = "Ablekuma South 2026 defense demonstration data"

FIELD_OFFICERS = (
    ("abs_field_01", "Adwoa", "Mensah", "0248101001"),
    ("abs_field_02", "Kwame", "Asare", "0248101002"),
    ("abs_field_03", "Esi", "Nyarko", "0248101003"),
    ("abs_field_04", "Kojo", "Tetteh", "0248101004"),
)

FIRST_NAMES = (
    "Ama", "Kwame", "Akosua", "Kofi", "Adwoa", "Yaw", "Abena", "Kwaku",
    "Esi", "Kojo", "Efua", "Kwesi", "Nana", "Nii", "Naana", "Mawuli",
    "Selorm", "Dzifa", "Afi", "Elikem", "Yaa", "Fiifi", "Araba", "Ekow",
)

SURNAMES = (
    "Mensah", "Owusu", "Asare", "Ofori", "Tetteh", "Addo", "Agyeman",
    "Boateng", "Annan", "Quaye", "Lamptey", "Lartey", "Sackey", "Aryee",
    "Adjei", "Badu", "Nyarko", "Acquah", "Darko", "Nartey", "Otoo",
    "Koomson", "Amoah", "Amponsah", "Ansah", "Appiah", "Essien", "Forson",
    "Gyamfi", "Koranteng", "Kwakye", "Opoku", "Sarpong", "Tagoe", "Yeboah",
    "Aidoo", "Adu", "Agyei", "Baah", "Dartey", "Frimpong", "Kusi",
)

STREET_NAMES = (
    "Abossey Okai Road", "Old Winneba Road", "Korle Bu Road", "Sakumo Street",
    "Beach Road", "Odorkor Road", "Dansoman High Street", "Mamprobi Link",
    "Glefe Road", "Sempe Lane", "Chorkor Avenue", "Korle Gonno Street",
    "Akotoku Crescent", "Banana Inn Road", "Tsui Bleoo Lane",
)

LANDMARKS = (
    "Near the community park", "Opposite the pharmacy", "Behind the market",
    "Near the lorry station", "Adjacent to the basic school",
    "Close to the main junction", "Beside the fuel station",
    "Near the community centre", "Opposite the clinic", "Behind the police post",
)

PROPERTY_USE_COUNTS = {
    "residential": 62,
    "commercial": 18,
    "industrial": 8,
    "mixed_use": 6,
    "institutional": 6,
}

VALUE_RANGES = {
    "residential": (Decimal("90000"), Decimal("650000"), Decimal("5000")),
    "commercial": (Decimal("250000"), Decimal("1500000"), Decimal("10000")),
    "industrial": (Decimal("500000"), Decimal("2200000"), Decimal("25000")),
    "mixed_use": (Decimal("300000"), Decimal("1250000"), Decimal("10000")),
    "institutional": (Decimal("350000"), Decimal("1800000"), Decimal("25000")),
}

# Approximate community centres used only to make the demonstration GPS data
# plausible. A small deterministic offset is added per property.
AREA_CENTRES = {
    "Akotoku": (Decimal("5.540300"), Decimal("-0.275900")),
    "Amanfong": (Decimal("5.539000"), Decimal("-0.271000")),
    "Banana Inn": (Decimal("5.550700"), Decimal("-0.264600")),
    "Beach Road": (Decimal("5.532900"), Decimal("-0.258500")),
    "Chorkor": (Decimal("5.531500"), Decimal("-0.268000")),
    "Coastal Executive Area": (Decimal("5.535100"), Decimal("-0.253600")),
    "Fishing Harbour Area": (Decimal("5.529900"), Decimal("-0.272800")),
    "Gbegbeyise": (Decimal("5.545000"), Decimal("-0.294000")),
    "Glefe": (Decimal("5.541000"), Decimal("-0.303000")),
    "High-End Mamprobi Enclave": (Decimal("5.541800"), Decimal("-0.259000")),
    "Korle Bu": (Decimal("5.537500"), Decimal("-0.261000")),
    "Korle Gonno": (Decimal("5.534500"), Decimal("-0.250000")),
    "Korle Lagoon View": (Decimal("5.535800"), Decimal("-0.246800")),
    "Korlebu Staff Area": (Decimal("5.538800"), Decimal("-0.259000")),
    "Mamprobi": (Decimal("5.543500"), Decimal("-0.263500")),
    "New Mamprobi": (Decimal("5.546500"), Decimal("-0.267000")),
    "Official Quarters": (Decimal("5.539700"), Decimal("-0.255000")),
    "Old Dansoman Extension": (Decimal("5.553000"), Decimal("-0.286000")),
    "Old Winneba Road Area": (Decimal("5.548500"), Decimal("-0.277000")),
    "Sempe": (Decimal("5.536500"), Decimal("-0.274500")),
    "South Chorkor": (Decimal("5.528700"), Decimal("-0.269500")),
    "South Korle Gonno Estates": (Decimal("5.532000"), Decimal("-0.252000")),
    "Tsui Bleoo": (Decimal("5.543000"), Decimal("-0.282000")),
    "VIP Residential Enclave": (Decimal("5.542500"), Decimal("-0.255500")),
}


class Command(BaseCommand):
    help = (
        "Replace Ablekuma South operational demo data with exactly 100 "
        "realistic properties and export a complete defense report."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Show the exact target and replacement plan without changing data.",
        )
        parser.add_argument(
            "--yes",
            action="store_true",
            help="Skip the typed confirmation (intended for controlled automation).",
        )
        parser.add_argument(
            "--output-dir",
            default="demo_exports",
            help="Directory for the XLSX report and credential reference.",
        )

    def handle(self, *args, **options):
        assembly = self._get_target_assembly()
        self._validate_configuration(assembly)

        before = self._assembly_snapshots()
        target_before = before[assembly.pk]
        self._print_plan(assembly, target_before, options["output_dir"])

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run only. No data was changed."))
            return

        if not options["yes"]:
            answer = input(
                f"\nType {CONFIRMATION_TEXT!r} to continue: "
            ).strip()
            if answer != CONFIRMATION_TEXT:
                raise CommandError("Confirmation did not match. No data was changed.")

        output_dir = self._resolve_output_dir(options["output_dir"])
        now = timezone.now()
        seed_result = None

        with transaction.atomic():
            self._clear_target_operational_data(assembly)
            assembly.active_billing_year = BILLING_YEAR
            assembly.save(update_fields=["active_billing_year", "updated_at"])
            seed_result = self._seed(assembly, now)
            self._validate_seed(assembly, seed_result)
            self._validate_other_assemblies_unchanged(before, assembly.pk)

        report_path, credentials_path = self._export_reports(
            assembly,
            seed_result,
            output_dir,
        )

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS("Ablekuma South demo population completed."))
        self.stdout.write(f"Properties: {PROPERTY_COUNT} (60 verified, 40 pending)")
        self.stdout.write("Verified bills: 50 paid, 5 part paid, 5 unpaid")
        self.stdout.write(f"Citizen/owner accounts: {OWNER_COUNT}")
        self.stdout.write(f"Field officer accounts: {len(FIELD_OFFICERS)}")
        self.stdout.write(f"Workbook: {report_path}")
        self.stdout.write(f"Login reference: {credentials_path}")

    def _get_target_assembly(self):
        candidates = Assembly.objects.filter(
            name__iexact=TARGET_ASSEMBLY_NAME,
            short_code__iexact=TARGET_ASSEMBLY_CODE,
        )
        count = candidates.count()
        if count != 1:
            raise CommandError(
                "Safety check failed: expected exactly one active assembly named "
                f"{TARGET_ASSEMBLY_NAME!r} with code {TARGET_ASSEMBLY_CODE!r}; "
                f"found {count}."
            )
        assembly = candidates.get()
        if not assembly.is_active:
            raise CommandError("Ablekuma South exists but is inactive.")
        return assembly

    def _validate_configuration(self, assembly):
        areas = CommunityArea.objects.filter(
            assembly=assembly,
            sub_metro=TARGET_ASSEMBLY_CODE,
            is_active=True,
        )
        if not areas.exists():
            raise CommandError("No active Ablekuma South community areas are configured.")

        policies = RatePolicy.objects.filter(
            assembly=assembly,
            year=BILLING_YEAR,
            status="active",
            is_active=True,
            covered_areas__in=areas,
        ).distinct()

        missing = [
            property_use
            for property_use in PROPERTY_USE_COUNTS
            if not policies.filter(property_use=property_use).exists()
        ]
        if missing:
            raise CommandError(
                "Missing active 2026 Ablekuma rate policies for: "
                + ", ".join(missing)
            )

    def _snapshot_for(self, assembly):
        return {
            "properties": Property.objects.filter(assembly=assembly).count(),
            "owners": Owner.objects.filter(assembly=assembly).count(),
            "bills": Bill.objects.filter(assembly=assembly).count(),
            "payments": Payment.objects.filter(assembly=assembly).count(),
            "assignments": FieldAssignment.objects.filter(assembly=assembly).count(),
            "visits": FieldVisit.objects.filter(assembly=assembly).count(),
            "enumerations": EnumerationRecord.objects.filter(assembly=assembly).count(),
        }

    def _assembly_snapshots(self):
        return {
            assembly.pk: self._snapshot_for(assembly)
            for assembly in Assembly.objects.order_by("pk")
        }

    def _print_plan(self, assembly, target_before, output_dir):
        self.stdout.write(self.style.WARNING("ABLEKUMA SOUTH REPLACEMENT PLAN"))
        self.stdout.write(
            f"Target: ID {assembly.pk} | {assembly.name} | {assembly.short_code}"
        )
        self.stdout.write(
            "Existing target records: "
            + ", ".join(f"{key}={value}" for key, value in target_before.items())
        )
        self.stdout.write(
            "This will replace Ablekuma South properties, owners/citizen links, "
            "bills, payments, field visits, assignments, enumerations and its "
            "field-officer accounts."
        )
        self.stdout.write(
            "It preserves assemblies, community areas, property classes, rate "
            "policies and administrative accounts."
        )
        self.stdout.write(
            "Okaikwai South and Ashiedu Keteke are excluded and will be verified unchanged."
        )
        self.stdout.write(
            "New target: 100 properties (60 verified / 40 pending), with verified "
            "bills split 50 paid / 5 part paid / 5 unpaid."
        )
        self.stdout.write(f"Reports directory: {output_dir}")

    def _resolve_output_dir(self, raw_path):
        output_dir = Path(raw_path)
        if not output_dir.is_absolute():
            output_dir = Path(settings.BASE_DIR) / output_dir
        output_dir.mkdir(parents=True, exist_ok=True)
        return output_dir.resolve()

    def _clear_target_operational_data(self, assembly):
        owner_user_ids = set(
            Owner.objects.filter(assembly=assembly)
            .exclude(linked_user_id=None)
            .values_list("linked_user_id", flat=True)
        )
        field_user_ids = set(
            UserProfile.objects.filter(
                assembly=assembly,
                role="field_officer",
            ).values_list("user_id", flat=True)
        )

        Payment.objects.filter(assembly=assembly).delete()
        FieldVisit.objects.filter(assembly=assembly).delete()
        FieldAssignment.objects.filter(assembly=assembly).delete()
        EnumerationRecord.objects.filter(assembly=assembly).delete()
        Bill.objects.filter(assembly=assembly).delete()
        Property.objects.filter(assembly=assembly).delete()
        Owner.objects.filter(assembly=assembly).delete()

        # Remove only target field officers and citizens that no longer own any
        # property in another assembly. Administrative accounts are preserved.
        User.objects.filter(pk__in=field_user_ids).delete()
        removable_citizen_ids = []
        for user in User.objects.filter(pk__in=owner_user_ids).select_related("profile"):
            profile = getattr(user, "profile", None)
            if (
                profile
                and profile.role == "citizen"
                and not Owner.objects.filter(linked_user=user).exists()
            ):
                removable_citizen_ids.append(user.pk)
        User.objects.filter(pk__in=removable_citizen_ids).delete()

    def _seed(self, assembly, now):
        rng = random.Random(20260916)
        administrator = self._registration_administrator(assembly)
        field_rows = self._create_field_officers(assembly)
        owner_rows = self._create_owners_and_accounts(assembly)
        property_rows = self._create_properties(
            assembly,
            administrator,
            field_rows,
            owner_rows,
            now,
            rng,
        )
        return {
            "administrator": administrator,
            "fields": field_rows,
            "owners": owner_rows,
            "properties": property_rows,
        }

    def _registration_administrator(self, assembly):
        profile = (
            UserProfile.objects.filter(
                assembly=assembly,
                role__in=("sys_admin", "revenue_admin", "property_officer"),
                user__is_active=True,
            )
            .select_related("user")
            .order_by("role", "user__username")
            .first()
        )
        if profile:
            return profile.user
        return User.objects.filter(is_superuser=True, is_active=True).order_by("pk").first()

    def _create_field_officers(self, assembly):
        rows = []
        for index, (username, first_name, last_name, phone) in enumerate(
            FIELD_OFFICERS,
            start=1,
        ):
            password = f"AbsField#{BILLING_YEAR}-{index:02d}"
            user = User.objects.create_user(
                username=username,
                password=password,
                first_name=first_name,
                last_name=last_name,
                email=f"{username}@demo.prms.local",
                is_active=True,
                is_staff=False,
            )
            profile = user.profile
            profile.role = "field_officer"
            profile.assembly = assembly
            profile.phone_number = phone
            profile.is_verified = True
            profile.must_change_password = False
            profile.employee_id = ""
            profile.save()
            profile.allowed_assemblies.add(assembly)
            rows.append(
                {
                    "user": user,
                    "profile": profile,
                    "password": password,
                    "phone": phone,
                }
            )
        return rows

    def _create_owners_and_accounts(self, assembly):
        rows = []
        for index in range(1, OWNER_COUNT + 1):
            first_name = FIRST_NAMES[(index - 1) % len(FIRST_NAMES)]
            surname = SURNAMES[(index * 7 - 1) % len(SURNAMES)]
            middle_initial = chr(65 + ((index * 5) % 26))
            full_name = f"{first_name} {middle_initial}. {surname}"
            network = ("024", "020", "055", "050", "054", "027")[index % 6]
            phone = f"{network}{6100000 + index:07d}"
            alternate = "" if index % 4 else f"026{7200000 + index:07d}"
            email = "" if index % 3 == 0 else f"abs.owner{index:03d}@demo.prms.local"
            ghana_card = f"GHA-{7300000000 + index:010d}-{index % 10}"
            owner = Owner.objects.create(
                assembly=assembly,
                full_name=full_name,
                primary_phone=phone,
                alternate_phone=alternate,
                email=email,
                ghana_card_number=ghana_card,
                postal_address=f"P.O. Box AS {100 + index}, Accra",
                residential_address=(
                    f"House No. {10 + index}, Ablekuma South, Greater Accra"
                ),
            )
            password = f"AbsCitizen#{index:03d}"
            account = link_or_create_citizen_for_owner(
                owner,
                default_password=password,
            )
            if not account["created"]:
                raise CommandError(
                    f"Safety check failed: citizen account was not newly created for {full_name}."
                )
            profile = account["user"].profile
            profile.is_verified = True
            profile.must_change_password = False
            profile.save(update_fields=["is_verified", "must_change_password", "updated_at"])
            rows.append(
                {
                    "owner": owner,
                    "user": account["user"],
                    "username": account["username"],
                    "password": password,
                }
            )
        return rows

    def _create_properties(
        self,
        assembly,
        administrator,
        field_rows,
        owner_rows,
        now,
        rng,
    ):
        policies_by_use = defaultdict(list)
        policies = (
            RatePolicy.objects.filter(
                assembly=assembly,
                year=BILLING_YEAR,
                status="active",
                is_active=True,
                covered_areas__sub_metro=TARGET_ASSEMBLY_CODE,
                covered_areas__is_active=True,
            )
            .prefetch_related("covered_areas")
            .distinct()
            .order_by("property_use", "property_class", "building_type")
        )
        for policy in policies:
            policies_by_use[policy.property_use].append(policy)

        owner_indices = list(range(OWNER_COUNT)) + list(range(28))
        rng.shuffle(owner_indices)

        use_slots = []
        for property_use, count in PROPERTY_USE_COUNTS.items():
            use_slots.extend([property_use] * count)
        rng.shuffle(use_slots)

        verification_slots = ["verified"] * VERIFIED_COUNT + ["pending"] * PENDING_COUNT
        rng.shuffle(verification_slots)
        payment_slots = (
            ["fully_paid"] * FULLY_PAID_COUNT
            + ["part_paid"] * PART_PAID_COUNT
            + ["unpaid"] * UNPAID_COUNT
        )
        rng.shuffle(payment_slots)
        payment_slot_index = 0

        registration_start = date(BILLING_YEAR, 1, 6)
        registration_end = min(timezone.localdate(), date(BILLING_YEAR, 9, 10))
        rows = []

        for index in range(PROPERTY_COUNT):
            property_number = index + 1
            owner_row = owner_rows[owner_indices[index]]
            property_use = use_slots[index]
            available_policies = policies_by_use[property_use]
            policy = available_policies[index % len(available_policies)]
            areas = [
                area
                for area in policy.covered_areas.all()
                if area.sub_metro == TARGET_ASSEMBLY_CODE and area.is_active
            ]
            if not areas:
                raise CommandError(f"Policy {policy.pk} has no active Ablekuma South area.")
            community = areas[(index * 3) % len(areas)]
            latitude, longitude = self._coordinates(community.name, property_number, rng)
            rateable_value = self._rateable_value(property_use, rng)
            registration_date = registration_start + timedelta(
                days=rng.randrange((registration_end - registration_start).days + 1)
            )
            registered_at = self._aware_datetime(
                registration_date,
                8 + index % 9,
                (index * 7) % 60,
            )
            status = verification_slots[index]
            floors = self._floor_count(property_use, rng)
            units = self._unit_count(property_use, policy.building_type, rng)
            occupancy = rng.choices(
                ("occupied", "vacant", "under_construction"),
                weights=(84, 11, 5),
                k=1,
            )[0]
            house_number = 12 + ((property_number * 17) % 480)
            street = STREET_NAMES[index % len(STREET_NAMES)]
            digital_address = f"GA-{100 + ((index * 37) % 900):03d}-{1000 + ((index * 83) % 9000):04d}"

            property_obj = Property.objects.create(
                assembly=assembly,
                owner=owner_row["owner"],
                digital_address=digital_address,
                address_line=f"House No. {house_number}, {street}",
                area_town=community.name,
                suburb=community.name,
                landmark=LANDMARKS[index % len(LANDMARKS)],
                community_area=community,
                latitude=latitude,
                longitude=longitude,
                property_use=property_use,
                building_type=policy.building_type,
                property_class=policy.property_class,
                applied_rate_policy=policy,
                number_of_floors=floors,
                number_of_units=units,
                year_built=rng.randint(1978, 2025),
                occupancy_status=occupancy,
                rateable_value=rateable_value,
                is_exempt=False,
                registration_source="admin",
                verification_status="pending",
                registered_by=administrator,
                is_active=True,
            )
            Property.objects.filter(pk=property_obj.pk).update(
                created_at=registered_at,
                updated_at=registered_at,
            )

            officer_row = field_rows[index % len(field_rows)]
            assigned_at = min(registered_at + timedelta(days=1 + index % 4), now)
            assignment = FieldAssignment.objects.create(
                assembly=assembly,
                property=property_obj,
                officer=officer_row["user"],
                assigned_by=administrator,
                status="assigned",
                priority=("medium", "high", "low")[index % 3],
                due_date=(assigned_at + timedelta(days=10)).date(),
                notes=f"{DEMO_NOTE}; assignment {property_number:03d}",
                assigned_at=assigned_at,
            )
            FieldAssignment.objects.filter(pk=assignment.pk).update(
                created_at=assigned_at,
                updated_at=assigned_at,
            )

            visit = None
            verified_at = None
            # The normal registration signal creates a current-year bill for
            # every active property, including properties awaiting verification.
            bill = Bill.objects.filter(
                property_record=property_obj,
                billing_year=BILLING_YEAR,
            ).first()
            if not bill:
                raise CommandError(
                    f"Registration bill was not created for {property_obj.property_reference}."
                )
            payments = []
            payment_category = "awaiting_verification"

            if status == "verified":
                verified_at = min(
                    assigned_at + timedelta(days=2 + index % 8, hours=2),
                    now - timedelta(hours=2),
                )
                property_obj.verification_status = "verified"
                property_obj.last_verified_at = verified_at
                property_obj.save(update_fields=["verification_status", "last_verified_at", "updated_at"])
                Property.objects.filter(pk=property_obj.pk).update(updated_at=verified_at)

                assignment.status = "completed"
                assignment.completed_at = verified_at
                assignment.save(update_fields=["status", "completed_at", "updated_at"])
                FieldAssignment.objects.filter(pk=assignment.pk).update(updated_at=verified_at)

                visit = FieldVisit.objects.create(
                    assembly=assembly,
                    property=property_obj,
                    officer=officer_row["user"],
                    visit_type="verification",
                    submitted_verification_type="property_verification",
                    payment_status_seen="unknown",
                    outcome="compliant",
                    submitted_verification_status="verified",
                    amount_outstanding_snapshot=bill.balance,
                    visited_at=verified_at,
                    latitude=latitude + Decimal("0.000015"),
                    longitude=longitude - Decimal("0.000012"),
                    gps_accuracy_meters=Decimal(str(5 + index % 12)),
                    gps_distance_meters=Decimal("2.10"),
                    gps_validation_status="matched",
                    notes=(
                        "Owner present: Yes\n\nProperty details, occupancy and GPS "
                        "coordinates confirmed during field verification."
                    ),
                )
                FieldVisit.objects.filter(pk=visit.pk).update(
                    created_at=verified_at,
                    updated_at=verified_at,
                )

                bill, _created, reason = generate_bill_for_property(
                    property_obj,
                    billing_year=BILLING_YEAR,
                )
                if not bill:
                    raise CommandError(
                        f"Bill generation failed for {property_obj.property_reference}: {reason}"
                    )
                issued_at = min(verified_at + timedelta(days=1), now - timedelta(hours=1))
                Bill.objects.filter(pk=bill.pk).update(
                    issued_at=issued_at,
                    created_at=issued_at,
                    updated_at=issued_at,
                )
                bill.refresh_from_db()

                payment_category = payment_slots[payment_slot_index]
                payment_slot_index += 1
                payments = self._create_payments(
                    bill,
                    owner_row["owner"],
                    payment_category,
                    issued_at,
                    now,
                    property_number,
                )
                bill.refresh_from_db()
            else:
                assignment.status = "in_progress" if index % 4 == 0 else "assigned"
                assignment.save(update_fields=["status", "updated_at"])

            property_obj.refresh_from_db()
            rows.append(
                {
                    "property": property_obj,
                    "owner_row": owner_row,
                    "officer_row": officer_row,
                    "assignment": assignment,
                    "visit": visit,
                    "bill": bill,
                    "payments": payments,
                    "payment_category": payment_category,
                    "registered_at": registered_at,
                    "verified_at": verified_at,
                }
            )

        return rows

    def _coordinates(self, area_name, index, rng):
        centre = AREA_CENTRES.get(
            area_name,
            (Decimal("5.540000"), Decimal("-0.270000")),
        )
        lat_offset = Decimal(str(rng.randint(-35, 35))) / Decimal("100000")
        lng_offset = Decimal(str(rng.randint(-35, 35))) / Decimal("100000")
        return (
            (centre[0] + lat_offset).quantize(Decimal("0.000001")),
            (centre[1] + lng_offset).quantize(Decimal("0.000001")),
        )

    def _rateable_value(self, property_use, rng):
        low, high, step = VALUE_RANGES[property_use]
        steps = int((high - low) / step)
        return (low + step * rng.randint(0, steps)).quantize(Decimal("0.01"))

    def _floor_count(self, property_use, rng):
        if property_use == "industrial":
            return rng.choice((1, 1, 2))
        if property_use in {"commercial", "mixed_use", "institutional"}:
            return rng.randint(1, 4)
        return rng.choices((1, 2, 3, 4), weights=(42, 38, 15, 5), k=1)[0]

    def _unit_count(self, property_use, building_type, rng):
        if building_type == "apartment":
            return rng.randint(4, 16)
        if building_type in {"compound", "shop", "office", "other"}:
            return rng.randint(2, 8)
        if property_use == "institutional":
            return rng.randint(4, 20)
        return rng.randint(1, 4)

    def _aware_datetime(self, day, hour, minute):
        naive = datetime.combine(day, time(hour=hour, minute=minute))
        return timezone.make_aware(naive, timezone.get_current_timezone())

    def _create_payments(
        self,
        bill,
        owner,
        category,
        issued_at,
        now,
        property_number,
    ):
        if category == "unpaid":
            return []

        due = bill.amount_due
        if category == "part_paid":
            amounts = [(due * Decimal("0.55")).quantize(Decimal("0.01"))]
        elif property_number % 3 == 0:
            first_amount = (due * Decimal("0.40")).quantize(Decimal("0.01"))
            amounts = [first_amount, due - first_amount]
        else:
            amounts = [due]

        rows = []
        for installment, amount in enumerate(amounts, start=1):
            paid_at = min(
                issued_at + timedelta(days=3 + installment * 4),
                now - timedelta(minutes=installment),
            )
            payment = Payment.objects.create(
                assembly=bill.assembly,
                property=bill.property_record,
                bill=bill,
                payer_name=owner.full_name,
                payer_phone=owner.primary_phone,
                payer_email=owner.email,
                property_reference_entered=bill.property_record.property_reference,
                amount=amount,
                payment_method=("momo", "cash", "bank", "pos", "card")[
                    (property_number + installment) % 5
                ],
                provider="PRMS Demo Collection",
                transaction_reference=(
                    f"ABS-DEMO-{BILLING_YEAR}-{property_number:03d}-{installment}"
                ),
                external_reference=f"DEMO-{property_number:03d}-{installment}",
                status="successful",
                paid_at=paid_at,
                raw_payload={
                    "source": "ablekuma_south_demo_seed",
                    "installment": installment,
                },
            )
            Payment.objects.filter(pk=payment.pk).update(
                created_at=paid_at,
                updated_at=paid_at,
            )
            payment.refresh_from_db()
            if not payment.applied_to_bill:
                raise CommandError(
                    f"Payment {payment.transaction_reference} was not applied to its bill."
                )
            rows.append(payment)
        return rows

    def _validate_seed(self, assembly, result):
        property_qs = Property.objects.filter(assembly=assembly)
        if property_qs.count() != PROPERTY_COUNT:
            raise CommandError("Post-seed validation failed: property count is not 100.")
        status_counts = dict(
            property_qs.values_list("verification_status")
            .annotate(total=Count("id"))
        )
        if status_counts != {"pending": PENDING_COUNT, "verified": VERIFIED_COUNT}:
            raise CommandError(f"Post-seed verification split is invalid: {status_counts}")

        bills = Bill.objects.filter(assembly=assembly, billing_year=BILLING_YEAR)
        verified_bills = bills.filter(property_record__verification_status="verified")
        verified_bill_counts = dict(
            verified_bills.values_list("status").annotate(total=Count("id"))
        )
        expected_verified_bills = {
            "paid": FULLY_PAID_COUNT,
            "part_paid": PART_PAID_COUNT,
            "pending": UNPAID_COUNT,
        }
        if verified_bill_counts != expected_verified_bills:
            raise CommandError(
                "Post-seed verified-property bill split is invalid: "
                f"{verified_bill_counts}"
            )
        if verified_bills.count() != VERIFIED_COUNT:
            raise CommandError("Post-seed validation failed: expected 60 verified bills.")
        if bills.count() != PROPERTY_COUNT:
            raise CommandError(
                "Post-seed validation failed: the registration workflow should "
                "have created one bill for each of the 100 properties."
            )
        pending_property_bills = bills.filter(
            property_record__verification_status="pending",
            status="pending",
            amount_paid=Decimal("0.00"),
        ).count()
        if pending_property_bills != PENDING_COUNT:
            raise CommandError(
                "Post-seed validation failed: expected 40 unpaid bills linked "
                "to properties awaiting verification."
            )
        if Owner.objects.filter(assembly=assembly).count() != OWNER_COUNT:
            raise CommandError("Post-seed validation failed: expected 72 owners.")
        if FieldAssignment.objects.filter(assembly=assembly).count() != PROPERTY_COUNT:
            raise CommandError("Post-seed validation failed: expected 100 assignments.")
        if FieldVisit.objects.filter(assembly=assembly).count() != VERIFIED_COUNT:
            raise CommandError("Post-seed validation failed: expected 60 field visits.")

        demo_field_ids = [row["user"].pk for row in result["fields"]]
        if UserProfile.objects.filter(
            user_id__in=demo_field_ids,
            assembly=assembly,
            role="field_officer",
        ).count() != len(FIELD_OFFICERS):
            raise CommandError("Post-seed validation failed: field officer access is invalid.")

        for owner in Owner.objects.filter(assembly=assembly).select_related("linked_user"):
            if not owner.linked_user_id or owner.linked_user.profile.role != "citizen":
                raise CommandError(f"Owner {owner.pk} is missing a valid citizen account.")

    def _validate_other_assemblies_unchanged(self, before, target_id):
        for assembly in Assembly.objects.exclude(pk=target_id):
            current = self._snapshot_for(assembly)
            if current != before[assembly.pk]:
                raise CommandError(
                    f"Safety validation failed: {assembly.name} changed from "
                    f"{before[assembly.pk]} to {current}. The transaction was rolled back."
                )

    def _export_reports(self, assembly, result, output_dir):
        timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
        report_path = output_dir / f"ablekuma_south_complete_report_{timestamp}.xlsx"
        credentials_path = output_dir / f"ablekuma_south_login_credentials_{timestamp}.txt"

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        self._write_summary(summary_sheet, assembly)
        self._write_properties_sheet(workbook, result["properties"])
        self._write_owners_sheet(workbook, result["owners"])
        self._write_field_sheet(workbook, result["fields"])
        self._write_assignments_sheet(workbook, result["properties"])
        self._write_bills_sheet(workbook, result["properties"])
        self._write_payments_sheet(workbook, result["properties"])

        for sheet in workbook.worksheets:
            self._format_sheet(sheet)
        workbook.save(report_path)
        self._write_credentials(credentials_path, assembly, result)
        return report_path, credentials_path

    def _write_summary(self, sheet, assembly):
        properties = Property.objects.filter(assembly=assembly)
        bills = Bill.objects.filter(assembly=assembly, billing_year=BILLING_YEAR)
        payments = Payment.objects.filter(assembly=assembly, status="successful")
        totals = bills.aggregate(total_due=Sum("amount_due"), total_paid=Sum("amount_paid"))
        due = totals["total_due"] or Decimal("0.00")
        paid = totals["total_paid"] or Decimal("0.00")
        rows = (
            ("ABLEKUMA SOUTH DEMONSTRATION DATA REPORT", ""),
            ("Generated", timezone.localtime().strftime("%Y-%m-%d %H:%M:%S %Z")),
            ("Assembly", f"{assembly.name} ({assembly.short_code})"),
            ("Billing year", BILLING_YEAR),
            ("Total properties", properties.count()),
            ("Verified properties", properties.filter(verification_status="verified").count()),
            ("Pending properties", properties.filter(verification_status="pending").count()),
            ("Owners / citizen accounts", Owner.objects.filter(assembly=assembly).count()),
            ("Field officers", UserProfile.objects.filter(assembly=assembly, role="field_officer").count()),
            ("Field assignments", FieldAssignment.objects.filter(assembly=assembly).count()),
            ("Completed field visits", FieldVisit.objects.filter(assembly=assembly).count()),
            ("Bills generated", bills.count()),
            (
                "Verified properties - fully paid",
                bills.filter(property_record__verification_status="verified", status="paid").count(),
            ),
            (
                "Verified properties - part paid",
                bills.filter(property_record__verification_status="verified", status="part_paid").count(),
            ),
            (
                "Verified properties - unpaid / outstanding",
                bills.filter(property_record__verification_status="verified", status="pending").count(),
            ),
            (
                "Pending properties - bills awaiting verification",
                bills.filter(property_record__verification_status="pending", status="pending").count(),
            ),
            ("Payment records", payments.count()),
            ("Total billed (GH\u00a2)", due),
            ("Revenue collected (GH\u00a2)", paid),
            ("Outstanding amount (GH\u00a2)", due - paid),
            ("Scope note", "Only Ablekuma South operational/demo records were replaced."),
        )
        for row in rows:
            sheet.append(row)
        sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="17324D")
        sheet["B1"].fill = PatternFill("solid", fgColor="17324D")

    def _write_properties_sheet(self, workbook, property_rows):
        sheet = workbook.create_sheet("Properties")
        sheet.append(
            (
                "Property Reference", "Owner", "Citizen Username", "Linked Property Count",
                "Address", "Community / Area", "Digital Address", "Latitude", "Longitude",
                "Property Use", "Building Type", "Property Class", "Rateable Value (GH\u00a2)",
                "Floors", "Units", "Year Built", "Occupancy", "Registered At",
                "Verification Status", "Verified At", "Field Officer", "Assignment Status",
                "Bill Number", "Bill Status", "Amount Due (GH\u00a2)", "Amount Paid (GH\u00a2)",
                "Outstanding (GH\u00a2)",
            )
        )
        owner_counts = dict(
            Property.objects.values_list("owner_id").annotate(total=Count("id"))
        )
        for row in sorted(property_rows, key=lambda item: item["property"].property_reference):
            prop = row["property"]
            bill = row["bill"]
            sheet.append(
                (
                    prop.property_reference,
                    prop.owner.full_name,
                    row["owner_row"]["username"],
                    owner_counts[prop.owner_id],
                    prop.address_line,
                    prop.community_area.name,
                    prop.digital_address,
                    prop.latitude,
                    prop.longitude,
                    prop.get_property_use_display(),
                    prop.get_building_type_display(),
                    prop.property_class,
                    prop.rateable_value,
                    prop.number_of_floors,
                    prop.number_of_units,
                    prop.year_built,
                    prop.get_occupancy_status_display(),
                    row["registered_at"].replace(tzinfo=None),
                    prop.get_verification_status_display(),
                    row["verified_at"].replace(tzinfo=None) if row["verified_at"] else "",
                    row["officer_row"]["user"].get_full_name(),
                    row["assignment"].get_status_display(),
                    bill.bill_number if bill else "",
                    bill.get_status_display() if bill else "Not billed",
                    bill.amount_due if bill else Decimal("0.00"),
                    bill.amount_paid if bill else Decimal("0.00"),
                    bill.balance if bill else Decimal("0.00"),
                )
            )

    def _write_owners_sheet(self, workbook, owner_rows):
        sheet = workbook.create_sheet("Owners & Citizen Logins")
        sheet.append(
            (
                "Owner", "Primary Phone / Login", "Alternate Phone", "Email", "Ghana Card",
                "Citizen Username", "Demo Password", "Properties Owned", "Linked Properties",
                "Account Role", "First-login Change Required",
            )
        )
        for row in owner_rows:
            owner = row["owner"]
            refs = list(
                owner.properties.order_by("property_reference")
                .values_list("property_reference", flat=True)
            )
            sheet.append(
                (
                    owner.full_name,
                    owner.primary_phone,
                    owner.alternate_phone,
                    owner.email,
                    owner.ghana_card_number,
                    row["username"],
                    row["password"],
                    len(refs),
                    ", ".join(refs),
                    "Citizen",
                    "No (demo account)",
                )
            )

    def _write_field_sheet(self, workbook, field_rows):
        sheet = workbook.create_sheet("Field Officer Logins")
        sheet.append(
            (
                "Officer", "Username", "Demo Password", "Phone", "Employee ID", "Role",
                "Total Assignments", "Completed", "Active / Pending", "Verified Properties",
            )
        )
        for row in field_rows:
            user = row["user"]
            assignments = FieldAssignment.objects.filter(officer=user)
            sheet.append(
                (
                    user.get_full_name(),
                    user.username,
                    row["password"],
                    row["phone"],
                    row["profile"].employee_id,
                    "Field Officer",
                    assignments.count(),
                    assignments.filter(status="completed").count(),
                    assignments.filter(status__in=("assigned", "in_progress")).count(),
                    assignments.filter(property__verification_status="verified").count(),
                )
            )

    def _write_assignments_sheet(self, workbook, property_rows):
        sheet = workbook.create_sheet("Assignments & Visits")
        sheet.append(
            (
                "Property Reference", "Field Officer", "Officer Username", "Priority",
                "Assigned At", "Due Date", "Assignment Status", "Completed At", "Visit ID",
                "Visited At", "Submitted Status", "Outcome", "GPS Validation", "GPS Distance (m)",
                "GPS Accuracy (m)", "Visit Notes",
            )
        )
        for row in property_rows:
            assignment = row["assignment"]
            visit = row["visit"]
            sheet.append(
                (
                    row["property"].property_reference,
                    row["officer_row"]["user"].get_full_name(),
                    row["officer_row"]["user"].username,
                    assignment.get_priority_display(),
                    assignment.assigned_at.replace(tzinfo=None),
                    assignment.due_date,
                    assignment.get_status_display(),
                    assignment.completed_at.replace(tzinfo=None) if assignment.completed_at else "",
                    visit.pk if visit else "",
                    visit.visited_at.replace(tzinfo=None) if visit else "",
                    visit.get_submitted_verification_status_display() if visit else "Pending",
                    visit.get_outcome_display() if visit else "Awaiting visit",
                    visit.get_gps_validation_status_display() if visit else "Not captured",
                    visit.gps_distance_meters if visit else "",
                    visit.gps_accuracy_meters if visit else "",
                    visit.notes if visit else "Assignment awaiting field verification.",
                )
            )

    def _write_bills_sheet(self, workbook, property_rows):
        sheet = workbook.create_sheet("Bills")
        sheet.append(
            (
                "Bill Number", "Property Reference", "Owner", "Billing Year", "Property Class",
                "Rateable Value (GH\u00a2)", "Rate Impost", "Principal (GH\u00a2)",
                "Amount Due (GH\u00a2)", "Amount Paid (GH\u00a2)", "Outstanding (GH\u00a2)",
                "Status", "Issued At", "Due Date", "Payment Category", "Payment References",
            )
        )
        for row in property_rows:
            bill = row["bill"]
            if not bill:
                continue
            sheet.append(
                (
                    bill.bill_number,
                    row["property"].property_reference,
                    row["property"].owner.full_name,
                    bill.billing_year,
                    bill.property_class_snapshot,
                    bill.rateable_value_snapshot,
                    bill.rate_impost_snapshot,
                    bill.principal_amount,
                    bill.amount_due,
                    bill.amount_paid,
                    bill.balance,
                    bill.get_status_display(),
                    bill.issued_at.replace(tzinfo=None),
                    bill.due_date,
                    row["payment_category"].replace("_", " ").title(),
                    ", ".join(payment.transaction_reference for payment in row["payments"]),
                )
            )

    def _write_payments_sheet(self, workbook, property_rows):
        sheet = workbook.create_sheet("Payments")
        sheet.append(
            (
                "Transaction Reference", "Property Reference", "Bill Number", "Payer", "Phone",
                "Amount (GH\u00a2)", "Method", "Status", "Reconciliation", "Paid At",
                "Applied To Bill",
            )
        )
        payments = [payment for row in property_rows for payment in row["payments"]]
        for payment in sorted(payments, key=lambda item: item.paid_at):
            sheet.append(
                (
                    payment.transaction_reference,
                    payment.property.property_reference,
                    payment.bill.bill_number,
                    payment.payer_name,
                    payment.payer_phone,
                    payment.amount,
                    payment.get_payment_method_display(),
                    payment.get_status_display(),
                    payment.get_reconciliation_status_display(),
                    payment.paid_at.replace(tzinfo=None),
                    "Yes" if payment.applied_to_bill else "No",
                )
            )

    def _format_sheet(self, sheet):
        if sheet.title == "Summary":
            sheet.column_dimensions["A"].width = 34
            sheet.column_dimensions["B"].width = 66
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, 1).font = Font(bold=True)
            for row in range(2, sheet.max_row + 1):
                if isinstance(sheet.cell(row, 2).value, Decimal):
                    sheet.cell(row, 2).number_format = '#,##0.00'
            return

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="17324D")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=8) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, Decimal):
                    cell.number_format = '#,##0.00'

    def _write_credentials(self, path, assembly, result):
        lines = [
            "ABLEKUMA SOUTH DEFENSE LOGIN REFERENCE",
            "=" * 42,
            f"Assembly: {assembly.name} ({assembly.short_code})",
            "Portal: use the normal application login page.",
            "These are demonstration accounts. Do not use these passwords for real users.",
            "",
            "FIELD OFFICERS",
            "-" * 42,
        ]
        for row in result["fields"]:
            user = row["user"]
            assignments = FieldAssignment.objects.filter(officer=user)
            lines.extend(
                [
                    f"Name: {user.get_full_name()}",
                    f"Username: {user.username}",
                    f"Password: {row['password']}",
                    f"Employee ID: {row['profile'].employee_id}",
                    (
                        "Assignments: "
                        f"{assignments.count()} total / "
                        f"{assignments.filter(status='completed').count()} completed / "
                        f"{assignments.filter(status__in=('assigned', 'in_progress')).count()} active"
                    ),
                    "",
                ]
            )

        lines.extend(["CITIZENS / PROPERTY OWNERS", "-" * 42])
        for row in result["owners"]:
            owner = row["owner"]
            refs = ", ".join(
                owner.properties.order_by("property_reference")
                .values_list("property_reference", flat=True)
            )
            lines.extend(
                [
                    f"Owner: {owner.full_name}",
                    f"Login phone / username: {row['username']}",
                    f"Password: {row['password']}",
                    f"Linked properties ({owner.properties.count()}): {refs}",
                    "",
                ]
            )
        path.write_text("\n".join(lines), encoding="utf-8")
