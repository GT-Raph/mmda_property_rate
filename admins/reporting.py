from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.paginator import Paginator
from django.db.models import Count, DecimalField, ExpressionWrapper, F, Q, Sum, Value
from django.db.models.functions import Coalesce, ExtractMonth, ExtractQuarter, ExtractYear
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from billing.models import Bill, BillDispute, CommunityArea
from payments.models import Payment
from properties.models import Owner, Property


MONEY_FIELD = DecimalField(max_digits=14, decimal_places=2)
REPORT_TYPES = (
    ("overview", "Overview"),
    ("properties", "Properties"),
    ("verification", "Verification"),
    ("billing_payments", "Billing & Payments"),
    ("citizens", "Citizens / Owners"),
    ("disputes", "Disputes"),
)


def _money(value):
    return f"GH₵ {Decimal(value or 0):,.0f}"


def _number(value):
    return f"{int(value or 0):,}"


def _period_config(period, selected_year):
    if period == "yearly":
        years = list(range(selected_year - 4, selected_year + 1))
        return {
            "labels": [str(year) for year in years],
            "buckets": years,
            "extractor": ExtractYear,
            "start": date(years[0], 1, 1),
            "end": date(selected_year + 1, 1, 1),
            "label": f"{years[0]}–{selected_year}",
            "period_label": "Yearly",
        }
    if period == "quarterly":
        return {
            "labels": ["Q1", "Q2", "Q3", "Q4"],
            "buckets": [1, 2, 3, 4],
            "extractor": ExtractQuarter,
            "start": date(selected_year, 1, 1),
            "end": date(selected_year + 1, 1, 1),
            "label": f"Q1–Q4 {selected_year}",
            "period_label": "Quarterly",
        }
    return {
        "labels": [date(2000, month, 1).strftime("%b") for month in range(1, 13)],
        "buckets": list(range(1, 13)),
        "extractor": ExtractMonth,
        "start": date(selected_year, 1, 1),
        "end": date(selected_year + 1, 1, 1),
        "label": f"January–December {selected_year}",
        "period_label": "Monthly",
    }


def _in_period(queryset, field, period):
    return queryset.filter(
        **{
            f"{field}__date__gte": period["start"],
            f"{field}__date__lt": period["end"],
        }
    )


def _series(queryset, field, period, aggregation):
    rows = (
        _in_period(queryset, field, period)
        .annotate(bucket=period["extractor"](field))
        .values("bucket")
        .annotate(value=aggregation)
        .order_by("bucket")
    )
    values = {int(row["bucket"]): row["value"] for row in rows if row["bucket"] is not None}
    return [float(values.get(bucket, 0) or 0) for bucket in period["buckets"]]


def _total(queryset, expression):
    return queryset.aggregate(
        total=Coalesce(Sum(expression), Value(Decimal("0.00"), output_field=MONEY_FIELD))
    )["total"]


def _cell(value, css_class="", url=""):
    return {"text": str(value if value not in (None, "") else "—"), "class": css_class, "url": url}


def _local_date(value):
    if not value:
        return "—"
    if hasattr(value, "hour"):
        value = timezone.localtime(value)
    return value.strftime("%d %b %Y")


def _dataset(label, values, color):
    return {"label": label, "data": values, "backgroundColor": color, "borderColor": color}


def _status_class(status):
    if status in {"verified", "paid", "resolved", "successful", "active"}:
        return "status-success"
    if status in {"rejected", "failed", "cancelled", "flagged", "possible_duplicate"}:
        return "status-danger"
    return "status-warning"


def build_filtered_report(request, assembly, params=None):
    params = params or request.GET
    today = timezone.localdate()
    report_type = (params.get("report_type") or "overview").strip()
    if report_type not in dict(REPORT_TYPES):
        report_type = "overview"
    period_key = (params.get("period") or "monthly").strip()
    if period_key not in {"monthly", "quarterly", "yearly"}:
        period_key = "monthly"
    try:
        selected_year = int(params.get("year") or today.year)
    except (TypeError, ValueError):
        selected_year = today.year
    if selected_year < 2000 or selected_year > today.year + 1:
        selected_year = today.year
    period = _period_config(period_key, selected_year)

    community_id = (params.get("community_area") or "").strip()
    property_use = (params.get("property_use") or "").strip()
    property_class = (params.get("property_class") or "").strip()
    building_type = (params.get("building_type") or "").strip()
    status = (params.get("status") or "").strip()
    payment_activity = (params.get("payment_activity") or "").strip()
    search = (params.get("search") or params.get("location") or "").strip()

    properties = Property.objects.select_related("owner", "community_area", "assembly")
    owners = Owner.objects.select_related("linked_user")
    bills = Bill.objects.select_related("property_record__owner", "property_record__community_area")
    payments = Payment.objects.filter(status="successful", paid_at__isnull=False).select_related(
        "property__owner", "bill__property_record__owner"
    )
    disputes = BillDispute.objects.select_related("bill__property_record__owner", "reviewed_by")
    if assembly:
        properties = properties.filter(assembly=assembly)
        owners = owners.filter(assembly=assembly)
        bills = bills.filter(assembly=assembly)
        payments = payments.filter(assembly=assembly)
        disputes = disputes.filter(bill__assembly=assembly)

    property_filters = Q()
    if community_id.isdigit():
        property_filters &= Q(community_area_id=int(community_id))
    if property_use:
        property_filters &= Q(property_use=property_use)
    if property_class:
        property_filters &= Q(property_class__icontains=property_class)
    if building_type:
        property_filters &= Q(building_type=building_type)
    if search and report_type in {"overview", "properties", "verification", "billing_payments"}:
        property_filters &= (
            Q(property_reference__icontains=search)
            | Q(owner__full_name__icontains=search)
            | Q(area_town__icontains=search)
            | Q(suburb__icontains=search)
            | Q(address_line__icontains=search)
            | Q(digital_address__icontains=search)
            | Q(community_area__name__icontains=search)
        )
    properties = properties.filter(property_filters)
    property_ids = properties.values("pk")
    bills = bills.filter(property_record_id__in=property_ids)
    payments = payments.filter(
        Q(property_id__in=property_ids) | Q(bill__property_record_id__in=property_ids)
    ).distinct()
    disputes = disputes.filter(bill__property_record_id__in=property_ids)

    if report_type == "citizens":
        if search:
            owners = owners.filter(
                Q(full_name__icontains=search)
                | Q(primary_phone__icontains=search)
                | Q(email__icontains=search)
            )
        if status == "linked":
            owners = owners.filter(linked_user__isnull=False)
        elif status == "unlinked":
            owners = owners.filter(linked_user__isnull=True)
        owner_ids = owners.values("pk")
        properties = properties.filter(owner_id__in=owner_ids)
        bills = bills.filter(property_record__owner_id__in=owner_ids)
        payments = payments.filter(
            Q(property__owner_id__in=owner_ids) | Q(bill__property_record__owner_id__in=owner_ids)
        ).distinct()
    elif report_type in {"properties", "verification"} and status:
        if status == "awaiting_review":
            properties = properties.filter(verification_status__in=["pending", "under_review"])
        else:
            properties = properties.filter(verification_status=status)
    elif report_type == "billing_payments" and status:
        if status == "outstanding":
            bills = bills.exclude(status="cancelled").filter(amount_due__gt=F("amount_paid"))
        else:
            bills = bills.filter(status=status)
    elif report_type == "disputes" and status:
        disputes = disputes.filter(status=status)

    if payment_activity == "recorded":
        paid_bill_ids = payments.exclude(bill_id=None).values("bill_id")
        bills = bills.filter(pk__in=paid_bill_ids)
    elif payment_activity == "none":
        paid_bill_ids = payments.exclude(bill_id=None).values("bill_id")
        bills = bills.exclude(pk__in=paid_bill_ids)

    period_properties = _in_period(properties, "created_at", period)
    period_owners = _in_period(owners, "created_at", period)
    period_bills = _in_period(bills, "issued_at", period)
    period_payments = _in_period(payments, "paid_at", period)
    period_disputes = _in_period(disputes, "created_at", period)
    financial_bills = period_bills.exclude(status="cancelled")
    outstanding_bills = financial_bills.filter(amount_due__gt=F("amount_paid"))
    balance = ExpressionWrapper(F("amount_due") - F("amount_paid"), output_field=MONEY_FIELD)

    def stat(title, value, scope, icon="property", icon_class="icon-blue"):
        return {"title": title, "value": value, "meta": scope, "icon_key": icon, "icon_class": icon_class}

    rows = []
    export_rows = []
    chart_is_money = False
    range_scope = period["label"]
    if report_type == "properties":
        verified = period_properties.filter(verification_status="verified")
        pending = period_properties.filter(verification_status__in=["pending", "under_review"])
        stats = [
            stat("Properties Registered", _number(period_properties.count()), range_scope),
            stat("Active Properties", _number(period_properties.filter(is_active=True).count()), "registered in scope", "mark", "icon-green"),
            stat("Verified", _number(verified.count()), "registered in scope", "mark", "icon-green"),
            stat("Pending", _number(pending.count()), "registered in scope", "alert", "icon-gold"),
            stat("Rateable Value", _money(_total(period_properties, "rateable_value")), "registered in scope", "money", "icon-gold"),
        ]
        chart_title = f"Property Registrations — {period['period_label']}, {selected_year}"
        chart_subtitle = "New property records matching the active filters"
        chart_sets = [_dataset("Properties Registered", _series(properties, "created_at", period, Count("id")), "#c9a031")]
        record_title, record_subtitle = "Property Records", "Properties registered within the selected scope"
        columns = ["Registered", "Reference", "Owner", "Community / Location", "Use", "Verification", "Status"]
        records = period_properties.order_by("-created_at")
        for obj in records:
            location = obj.community_area.name if obj.community_area_id else (obj.area_town or obj.address_line)
            values = [_local_date(obj.created_at), obj.property_reference, obj.owner.full_name, location, obj.get_property_use_display(), obj.get_verification_status_display(), "Active" if obj.is_active else "Inactive"]
            rows.append({"cells": [_cell(values[0]), _cell(values[1], url=reverse("properties:property-detail", args=[obj.pk])), _cell(values[2]), _cell(values[3]), _cell(values[4]), _cell(values[5], _status_class(obj.verification_status)), _cell(values[6], _status_class(values[6].lower()))]})
            export_rows.append(values)
    elif report_type == "verification":
        denominator = period_properties.count()
        verified = period_properties.filter(verification_status="verified")
        pending = period_properties.filter(verification_status__in=["pending", "under_review"])
        flagged = period_properties.filter(verification_status__in=["flagged", "possible_duplicate", "rejected"])
        rate = (verified.count() / denominator * 100) if denominator else 0
        stats = [
            stat("Properties in Scope", _number(denominator), range_scope),
            stat("Verified", _number(verified.count()), "same registration cohort", "mark", "icon-green"),
            stat("Pending", _number(pending.count()), "same registration cohort", "alert", "icon-gold"),
            stat("Flagged / Rejected", _number(flagged.count()), "same registration cohort", "alert", "icon-blue"),
            stat("Verification Rate", f"{rate:.1f}%", "verified ÷ properties in scope", "collection", "icon-green"),
        ]
        chart_title = f"Verification Status — {period['period_label']}, {selected_year}"
        chart_subtitle = "Current status of properties registered in each period"
        chart_sets = [
            _dataset("Verified", _series(properties.filter(verification_status="verified"), "created_at", period, Count("id")), "#16a34a"),
            _dataset("Pending / Review", _series(properties.filter(verification_status__in=["pending", "under_review"]), "created_at", period, Count("id")), "#d9a92c"),
            _dataset("Flagged / Rejected", _series(properties.filter(verification_status__in=["flagged", "possible_duplicate", "rejected"]), "created_at", period, Count("id")), "#dc3545"),
        ]
        record_title, record_subtitle = "Verification Records", "Verification state for the selected registration cohort"
        columns = ["Registered", "Property", "Owner", "Community / Location", "Status", "Last Verified"]
        records = period_properties.order_by("-created_at")
        for obj in records:
            location = obj.community_area.name if obj.community_area_id else obj.area_town
            values = [_local_date(obj.created_at), obj.property_reference, obj.owner.full_name, location, obj.get_verification_status_display(), _local_date(obj.last_verified_at)]
            rows.append({"cells": [_cell(values[0]), _cell(values[1], url=reverse("properties:property-detail", args=[obj.pk])), _cell(values[2]), _cell(values[3]), _cell(values[4], _status_class(obj.verification_status)), _cell(values[5])]})
            export_rows.append(values)
    elif report_type == "billing_payments":
        revenue = _total(period_payments, "amount")
        outstanding = _total(outstanding_bills, balance)
        stats = [
            stat("Bills Generated", _number(period_bills.count()), range_scope, "bill"),
            stat("Amount Billed", _money(_total(financial_bills, "amount_due")), range_scope, "money"),
            stat("Payments Recorded", _number(period_payments.count()), "successful only", "paid", "icon-green"),
            stat("Revenue Collected", _money(revenue), "successful payments in scope", "collection", "icon-green"),
            stat("Outstanding Amount", _money(outstanding), "due less paid for bills in scope", "alert", "icon-gold"),
        ]
        chart_is_money = True
        chart_title = f"Billing & Collections — {period['period_label']}, {selected_year}"
        chart_subtitle = "Amounts billed and successfully collected"
        chart_sets = [
            _dataset("Amount Billed", _series(bills.exclude(status="cancelled"), "issued_at", period, Sum("amount_due")), "#1f1f1f"),
            _dataset("Revenue Collected", _series(payments, "paid_at", period, Sum("amount")), "#16a34a"),
        ]
        record_title, record_subtitle = "Billing Records", "Bills issued within the selected scope"
        columns = ["Issued", "Bill", "Property", "Owner", "Due", "Paid", "Outstanding", "Status"]
        records = period_bills.order_by("-issued_at")
        for obj in records:
            remaining = max((obj.amount_due or 0) - (obj.amount_paid or 0), Decimal("0.00"))
            values = [_local_date(obj.issued_at), obj.bill_number, obj.property_record.property_reference, obj.property_record.owner.full_name, float(obj.amount_due or 0), float(obj.amount_paid or 0), float(remaining), obj.get_status_display()]
            rows.append({"cells": [_cell(values[0]), _cell(values[1]), _cell(values[2], url=reverse("properties:property-detail", args=[obj.property_record_id])), _cell(values[3]), _cell(_money(values[4])), _cell(_money(values[5])), _cell(_money(values[6])), _cell(values[7], _status_class(obj.status))]})
            export_rows.append(values)
    elif report_type == "citizens":
        owners_with_properties = period_owners.filter(properties__isnull=False).distinct()
        linked = period_owners.filter(linked_user__isnull=False)
        scoped_property_count = properties.filter(owner_id__in=period_owners.values("pk")).count()
        stats = [
            stat("Property Owners Added", _number(period_owners.count()), range_scope, "assembly"),
            stat("Owners With Properties", _number(owners_with_properties.count()), "owners added in scope", "property"),
            stat("Linked Portal Accounts", _number(linked.count()), "owners added in scope", "mark", "icon-green"),
            stat("Properties Owned", _number(scoped_property_count), "owned by owners added in scope", "property"),
        ]
        chart_title = f"Owner Registrations — {period['period_label']}, {selected_year}"
        chart_subtitle = "New Property Owner records and linked properties"
        chart_sets = [_dataset("Property Owners", _series(owners, "created_at", period, Count("id")), "#c9a031")]
        record_title, record_subtitle = "Property Owner Records", "Owner records created within the selected scope"
        columns = ["Registered", "Owner", "Phone", "Email", "Properties", "Portal Account"]
        records = period_owners.annotate(property_count=Count("properties", distinct=True)).order_by("-created_at")
        for obj in records:
            values = [_local_date(obj.created_at), obj.full_name, obj.primary_phone, obj.email or "Not provided", obj.property_count, "Linked" if obj.linked_user_id else "Not linked"]
            rows.append({"cells": [_cell(values[0]), _cell(values[1]), _cell(values[2]), _cell(values[3]), _cell(values[4]), _cell(values[5], "status-success" if obj.linked_user_id else "status-neutral")]})
            export_rows.append(values)
    elif report_type == "disputes":
        total = period_disputes.count()
        submitted = period_disputes.filter(status="submitted").count()
        review = period_disputes.filter(status="under_review").count()
        resolved = period_disputes.filter(status="resolved").count()
        stats = [
            stat("Disputes Opened", _number(total), range_scope, "alert"),
            stat("Submitted", _number(submitted), "in selected scope", "alert", "icon-gold"),
            stat("Under Review", _number(review), "in selected scope", "alert", "icon-blue"),
            stat("Resolved", _number(resolved), "in selected scope", "mark", "icon-green"),
            stat("Resolution Rate", f"{(resolved / total * 100) if total else 0:.1f}%", "resolved ÷ disputes opened", "collection", "icon-green"),
        ]
        chart_title = f"Disputes Opened — {period['period_label']}, {selected_year}"
        chart_subtitle = "Disputes created within each reporting period"
        chart_sets = [_dataset("Disputes", _series(disputes, "created_at", period, Count("id")), "#c9a031")]
        record_title, record_subtitle = "Dispute Records", "Disputes opened within the selected scope"
        columns = ["Date", "Bill", "Property / Owner", "Type", "Status", "Reviewed"]
        records = period_disputes.order_by("-created_at")
        for obj in records:
            owner_detail = f"{obj.bill.property_record.property_reference} · {obj.bill.property_record.owner.full_name}"
            values = [_local_date(obj.created_at), obj.bill.bill_number, owner_detail, obj.get_issue_type_display(), obj.get_status_display(), _local_date(obj.reviewed_at)]
            rows.append({"cells": [_cell(values[0]), _cell(values[1]), _cell(values[2]), _cell(values[3]), _cell(values[4], _status_class(obj.status)), _cell(values[5])]})
            export_rows.append(values)
    else:
        all_properties = properties
        verified = all_properties.filter(is_active=True, verification_status="verified")
        pending = all_properties.filter(is_active=True, verification_status__in=["pending", "under_review"])
        stats = [
            stat("Total Properties", _number(all_properties.count()), "as of today"),
            stat("Verified Properties", _number(verified.count()), "active, as of today", "mark", "icon-green"),
            stat("Pending Verification", _number(pending.count()), "active, as of today", "alert", "icon-gold"),
            stat("Revenue Collected", _money(_total(period_payments, "amount")), range_scope, "collection", "icon-green"),
            stat("Outstanding Amount", _money(_total(outstanding_bills, balance)), "bills issued in scope", "alert", "icon-gold"),
        ]
        chart_is_money = True
        chart_title = f"Revenue Collected — {period['period_label']}, {selected_year}"
        chart_subtitle = "Successful payments matching the active scope"
        chart_sets = [_dataset("Revenue Collected", _series(payments, "paid_at", period, Sum("amount")), "#16a34a")]
        record_title, record_subtitle = "Recent Financial Activity", "Successful payments received within the selected scope"
        columns = ["Date", "Transaction", "Property", "Payer", "Method", "Amount"]
        records = period_payments.order_by("-paid_at")
        for obj in records:
            prop = obj.property or (obj.bill.property_record if obj.bill_id else None)
            values = [_local_date(obj.paid_at), obj.transaction_reference, getattr(prop, "property_reference", obj.property_reference_entered or "Unmatched"), obj.payer_name or getattr(getattr(prop, "owner", None), "full_name", "—"), obj.get_payment_method_display(), float(obj.amount or 0)]
            rows.append({"cells": [_cell(values[0]), _cell(values[1]), _cell(values[2]), _cell(values[3]), _cell(values[4]), _cell(_money(values[5]))]})
            export_rows.append(values)

    paginator = Paginator(rows, 10)
    page = paginator.get_page(params.get("page", 1))
    query = params.copy()
    query.pop("page", None)
    query.pop("csrfmiddlewaretoken", None)
    query.pop("action", None)

    communities = CommunityArea.objects.filter(is_active=True)
    if assembly:
        communities = communities.filter(assembly=assembly)
    active_filters = [{"label": "Period", "value": range_scope}]
    labels = {
        "community_area": dict(communities.values_list("pk", "name")).get(int(community_id)) if community_id.isdigit() else "",
        "property_use": dict(Property.PROPERTY_USE_CHOICES).get(property_use, ""),
        "building_type": dict(Property.BUILDING_TYPE_CHOICES).get(building_type, ""),
        "property_class": property_class,
        "status": status.replace("_", " ").title(),
        "payment_activity": payment_activity.replace("_", " ").title(),
        "search": search,
    }
    filter_titles = {"community_area": "Community", "property_use": "Use", "building_type": "Building", "property_class": "Class", "status": "Status", "payment_activity": "Payment", "search": "Search"}
    for key, value in labels.items():
        if value:
            active_filters.append({"label": filter_titles[key], "value": value})

    status_choices = {
        "overview": (),
        "properties": Property.VERIFICATION_STATUS_CHOICES,
        "verification": (("awaiting_review", "Pending / Under Review"), *Property.VERIFICATION_STATUS_CHOICES),
        "billing_payments": (("outstanding", "Outstanding Balance"), *Bill.STATUS_CHOICES),
        "citizens": (("linked", "Portal Account Linked"), ("unlinked", "No Portal Account")),
        "disputes": BillDispute.STATUS_CHOICES,
    }[report_type]

    return {
        "report_type_choices": REPORT_TYPES,
        "selected_report_type": report_type,
        "selected_report_type_label": dict(REPORT_TYPES)[report_type],
        "selected_period": period_key,
        "selected_year": selected_year,
        "year_options": list(range(today.year + 1, today.year - 6, -1)),
        "community_area_choices": communities.order_by("name"),
        "selected_community_area": community_id,
        "property_use_choices": Property.PROPERTY_USE_CHOICES,
        "selected_property_use": property_use,
        "building_type_choices": Property.BUILDING_TYPE_CHOICES,
        "selected_building_type": building_type,
        "property_class": property_class,
        "search": search,
        "selected_status": status,
        "status_choices": status_choices,
        "selected_payment_activity": payment_activity,
        "active_filters": active_filters,
        "stats": stats,
        "range_label": range_scope,
        "chart_title": chart_title,
        "chart_subtitle": chart_subtitle,
        "chart_config": {"labels": period["labels"], "datasets": chart_sets, "isMoney": chart_is_money},
        "record_title": record_title,
        "record_subtitle": record_subtitle,
        "record_columns": columns,
        "record_rows": list(page.object_list),
        "record_page_obj": page,
        "record_paginator": paginator,
        "record_start_index": page.start_index() if paginator.count else 0,
        "record_end_index": page.end_index() if paginator.count else 0,
        "filter_query": query.urlencode(),
        "export_query_fields": [(key, value) for key, value in query.items()],
        "export_title": f"{dict(REPORT_TYPES)[report_type]} Report",
        "export_headers": columns,
        "export_rows": export_rows,
    }


def export_filtered_report(request, report):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Filtered Results"
    worksheet.append([report["export_title"]])
    worksheet.append(["Generated", timezone.localtime().strftime("%d %b %Y %H:%M")])
    worksheet.append(["Scope", " | ".join(f"{item['label']}: {item['value']}" for item in report["active_filters"])])
    worksheet.append([])
    worksheet.append(report["export_headers"])
    for row in report["export_rows"]:
        worksheet.append(row)

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(report["export_headers"]), 1))
    worksheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F1F1F")
    worksheet["A1"].alignment = Alignment(horizontal="left")
    for cell in worksheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8A6500")
    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        letter = get_column_letter(index)
        length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[letter].width = min(max(length + 2, 12), 42)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{report['selected_report_type']}_{report['selected_period']}_{report['selected_year']}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
