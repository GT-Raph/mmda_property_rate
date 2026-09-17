import json
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.core.management.base import CommandError
from django.db.models import F
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from admins.models import Assembly, UserProfile
from billing.models import Bill, BillDispute, CommunityArea, PropertyClass, RatePolicy
from field_operations.models import FieldAssignment, FieldVisit
from payments.models import Payment
from properties.models import Owner, Property


@override_settings(
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
)
class PopulateAblekumaSouthCommandTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.media_override = override_settings(MEDIA_ROOT=self.temp.name)
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.assembly = Assembly.objects.create(
            name="Ablekuma South", short_code="AMA", assembly_type="metro",
            region="Greater Accra", active_billing_year=timezone.localdate().year,
        )
        self.oks = CommunityArea.objects.create(
            assembly=self.assembly, sub_metro="OKS", name="Existing Okaikoi Area",
        )
        self.aks = CommunityArea.objects.create(
            assembly=self.assembly, sub_metro="AKS", name="Existing Ashiedu Keteke Area",
        )
        self.areas = [
            CommunityArea.objects.create(
                assembly=self.assembly, sub_metro="ABS", name=f"Approved Ablekuma Area {index}",
            ) for index in range(1, 9)
        ]
        for use in ["residential", "commercial", "industrial", "mixed_use", "institutional"]:
            property_class = PropertyClass.objects.create(
                assembly=self.assembly, name=f"{use.replace('_', ' ').title()} Class",
            )
            property_class.covered_areas.set(self.areas)
            policy = RatePolicy.objects.create(
                assembly=self.assembly, policy_name=f"Approved {use} policy",
                year=timezone.localdate().year, property_use=use,
                property_class=property_class.name, building_type="",
                rate_impost="0.01000", minimum_charge="100.00",
                effective_from=timezone.localdate().replace(month=1, day=1),
                effective_to=timezone.localdate().replace(month=12, day=31), status="active",
            )
            policy.covered_areas.set(self.areas)

    def command(self, *args, **kwargs):
        stdout, stderr = StringIO(), StringIO()
        call_command("populate_ablekuma_south", *args, stdout=stdout, stderr=stderr, **kwargs)
        return stdout.getvalue(), stderr.getvalue()

    def test_default_is_dry_run_and_preserves_other_sub_metros(self):
        stdout, _stderr = self.command(assembly_id=self.assembly.pk)
        self.assertIn("DRY-RUN ONLY", stdout)
        self.assertEqual(Property.objects.count(), 0)
        self.assertTrue(CommunityArea.objects.filter(pk=self.oks.pk, name="Existing Okaikoi Area").exists())
        self.assertTrue(CommunityArea.objects.filter(pk=self.aks.pk, name="Existing Ashiedu Keteke Area").exists())

    def test_confirm_creates_exact_coherent_population_and_private_report(self):
        self.command(assembly_id=self.assembly.pk, confirm=True)
        props = Property.objects.filter(assembly=self.assembly)
        verified = props.filter(verification_status="verified")
        pending = props.exclude(verification_status="verified")
        self.assertEqual(props.count(), 100)
        self.assertFalse(props.exclude(community_area__sub_metro="ABS").exists())
        self.assertEqual(verified.count(), 60)
        self.assertEqual(pending.count(), 40)
        self.assertEqual(Owner.objects.filter(assembly=self.assembly).count(), 85)
        self.assertEqual(UserProfile.objects.filter(role="citizen", assembly=self.assembly).count(), 85)
        self.assertEqual(UserProfile.objects.filter(role="field_officer", assembly=self.assembly).count(), 4)
        self.assertEqual(FieldAssignment.objects.filter(property__in=props).count(), 100)
        self.assertEqual(FieldVisit.objects.filter(property__in=props).count(), 60)
        self.assertEqual(FieldAssignment.objects.filter(property__in=verified, status="completed").count(), 60)
        bills = Bill.objects.filter(property_record__in=verified, billing_year=timezone.localdate().year)
        self.assertEqual(bills.filter(status="paid").count(), 50)
        self.assertEqual(bills.filter(status="part_paid").count(), 5)
        self.assertEqual(bills.filter(amount_paid=0).count(), 5)
        self.assertFalse(Payment.objects.filter(property__in=pending, status="successful", applied_to_bill=True).exists())
        generated_payments = Payment.objects.filter(property__in=verified)
        self.assertEqual(generated_payments.count(), 55)
        self.assertEqual(generated_payments.filter(payment_method="momo").count(), 55)
        self.assertFalse(generated_payments.exclude(provider__in=["mtn", "vodafone", "airteltigo"]).exists())
        self.assertFalse(generated_payments.exclude(external_reference=F("transaction_reference")).exists())
        for payment in generated_payments:
            self.assertLessEqual(payment.amount, payment.bill.amount_due)
            self.assertTrue(payment.applied_to_bill)
        self.assertEqual(BillDispute.objects.filter(bill__property_record__in=props).count(), 10)
        marker = Path(self.temp.name) / "private_population_reports" / f"ablekuma_south_assembly_{self.assembly.pk}.json"
        self.assertTrue(marker.exists())
        manifest = json.loads(marker.read_text(encoding="utf-8"))
        report = Path(manifest["report_path"])
        self.assertTrue(report.exists())
        workbook = load_workbook(report, read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "SUMMARY", "OWNERS", "CITIZEN CREDENTIALS",
                "FIELD OFFICER CREDENTIALS", "FIELD OFFICER ASSIGNMENTS",
                "PROPERTIES", "BILLS", "PAYMENTS", "FIELD VISITS",
                "BILL DISPUTES", "CITIZEN DISPUTES",
            ],
        )
        self.assertEqual(workbook["PROPERTIES"].max_row, 101)
        self.assertEqual(workbook["OWNERS"].max_row, 86)
        self.assertEqual(workbook["BILLS"].max_row, 101)
        self.assertEqual(workbook["PAYMENTS"].max_row, 56)
        self.assertEqual(workbook["CITIZEN DISPUTES"].max_row, 3)
        workbook.close()
        self.assertEqual(CommunityArea.objects.filter(pk=self.oks.pk).count(), 1)
        self.assertEqual(CommunityArea.objects.filter(pk=self.aks.pk).count(), 1)

    def test_duplicate_confirmed_execution_is_blocked(self):
        self.command(assembly_id=self.assembly.pk, confirm=True)
        with self.assertRaises(CommandError):
            self.command(assembly_id=self.assembly.pk, confirm=True)
        self.assertEqual(Property.objects.filter(assembly=self.assembly).count(), 100)

    def test_stale_manifest_allows_safe_repopulation(self):
        marker_dir = Path(self.temp.name) / "private_population_reports"
        marker_dir.mkdir(parents=True)
        marker = marker_dir / f"ablekuma_south_assembly_{self.assembly.pk}.json"
        marker.write_text(
            json.dumps(
                {
                    "status": "completed",
                    "execution_id": "stale-run",
                    "assembly_id": self.assembly.pk,
                    "owner_ids": list(range(990000, 990085)),
                    "user_ids": list(range(991000, 991085)),
                    "field_officer_user_ids": list(range(992000, 992004)),
                    "property_ids": list(range(993000, 993100)),
                    "bill_ids": list(range(994000, 994100)),
                    "payment_ids": list(range(995000, 995055)),
                    "assignment_ids": list(range(996000, 996100)),
                    "field_visit_ids": list(range(997000, 997060)),
                    "bill_dispute_ids": list(range(998000, 998010)),
                    "citizen_dispute_ids": list(range(999000, 999002)),
                }
            ),
            encoding="utf-8",
        )

        stdout, _stderr = self.command(assembly_id=self.assembly.pk, confirm=True)

        self.assertIn("population completed", stdout)
        self.assertEqual(Property.objects.filter(assembly=self.assembly).count(), 100)
        self.assertTrue(
            (marker_dir / f"ablekuma_south_assembly_{self.assembly.pk}_stale-run.json").exists()
        )

    def test_non_abs_master_data_is_rejected(self):
        CommunityArea.objects.filter(assembly=self.assembly, sub_metro="ABS").delete()
        with self.assertRaises(CommandError):
            self.command(assembly_id=self.assembly.pk, confirm=True)
        self.assertEqual(Property.objects.count(), 0)
